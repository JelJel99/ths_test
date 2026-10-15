import nest_asyncio
nest_asyncio.apply()

import os
import re
import time
import random
import asyncio
import urllib.parse
from datetime import datetime, date, timedelta
import pandas as pd
import requests
import httpx
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from playwright.async_api import async_playwright
from tqdm.auto import tqdm

# ==========================================
# GLOBAL CONFIG
# ==========================================
TARGET = 1000
PORTAL_CAP = 250
MAX_DATE = date(2026, 10, 3)
MIN_DATE = date(2025, 10, 3)
OUTPUT_COLUMNS = ["id", "source", "date", "title", "label", "url"]

USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:124.0) Gecko/20100101 Firefox/124.0"
]

MONTH_MAP = {
    "jan": "01", "feb": "02", "mar": "03", "apr": "04", "mei": "05", "jun": "06",
    "jul": "07", "agu": "08", "sep": "09", "okt": "10", "nov": "11", "des": "12"
}

DOMAINS = {
    "technology": {
        "label": "Technology",
        "output_file": "data/technology.csv",
        "keywords": ["software", "hardware", "it", "teknologi", "wifi", "internet", "robot", "ict"],
        "cnn_indeks": "teknologi/indeks/8",
        "cnn_section": "/teknologi/",
        "kompas_site": "tekno",
        "kompas_host": "tekno.kompas.com",
        "kompas_url_patterns": ["tekno.kompas.com/read/", "www.kompas.com/tekno/read/"],
        "detik_base": "https://inet.detik.com/indeks",
        "detik_host": "inet.detik.com",
        "detik_date_iso": True,
        "tribun_channels": ["index-news/techno"],
        "tribun_marker": "techno",
    },
    "health": {
        "label": "Health",
        "output_file": "data/health.csv",
        "keywords": ["kesehatan", "sehat", "bugar", "obat", "pertolongan", "pencegahan", "vitamin"],
        "cnn_indeks": "gaya-hidup/indeks/10",
        "cnn_section": "/gaya-hidup/",
        "kompas_site": "health",
        "kompas_host": "health.kompas.com",
        "kompas_url_patterns": ["health.kompas.com/read/", "www.kompas.com/health/read/"],
        "detik_base": "https://health.detik.com/indeks",
        "detik_host": "health.detik.com",
        "detik_date_iso": False,
        "tribun_channels": ["index-news/kesehatan"],
        "tribun_marker": "kesehatan",
    },
    "education": {
        "label": "Education",
        "output_file": "data/education.csv",
        "keywords": ["kuliah", "materi", "pendidikan", "kelas", "pembelajaran", "guru", "dosen"],
        "cnn_indeks": "edukasi/indeks/559",
        "cnn_section": "/edukasi/",
        "kompas_site": "edukasi",
        "kompas_host": "www.kompas.com/edu",
        "kompas_url_patterns": ["www.kompas.com/edu/read/", "edukasi.kompas.com/read/"],
        "detik_base": "https://www.detik.com/edu/indeks",
        "detik_host": "detik.com/edu",
        "detik_date_iso": False,
        "tribun_channels": ["pendidikan-go"],
        "tribun_marker": "pendidikan",
    },
    "sport": {
        "label": "Sport",
        "output_file": "data/sport.csv",
        "keywords": ["sepakbola", "bola", "soccer", "nbl", "basket", "volly", "asian games", "skor", "ht ft"],
        "cnn_indeks": "olahraga/indeks/7",
        "cnn_section": "/olahraga/",
        "kompas_site": "bola",
        "kompas_host": "bola.kompas.com",
        "kompas_url_patterns": ["bola.kompas.com/read/", "www.kompas.com/sports/read/"],
        "detik_base": "https://sport.detik.com/indeks",
        "detik_host": "sport.detik.com",
        "detik_date_iso": False,
        "tribun_channels": ["index-news/sport"],
        "tribun_marker": "sport",
    },
    "politics": {
        "label": "Politics",
        "output_file": "data/politic.csv",
        "keywords": ["politik", "parpol", "partai", "pemilu", "kpu", "pencoblosan", "pemungutan suara", "oposisi"],
        "cnn_indeks": "nasional/indeks/3",
        "cnn_section": "/nasional/",
        "kompas_site": "nasional",
        "kompas_host": "nasional.kompas.com",
        "kompas_url_patterns": [
            "nasional.kompas.com/read/",
            "www.kompas.com/nasional/read/",
            "sorotpolitik.kompas.com/read/",
        ],
        "detik_base": "https://news.detik.com/indeks",
        "detik_host": "news.detik.com",
        "detik_date_iso": False,
        "tribun_channels": [
            "index-news/nasional",
            "index-news/dpd-ri",
            "index-news/mpr-ri",
            "index-news/kilas-kementerian",
        ],
        "tribun_marker": "nasional",
    },
}

PORTAL_ORDER = ["Detik", "Kompas", "CNN Indonesia", "Tribun News"]


# ==========================================
# SHARED HELPERS (unchanged)
# ==========================================
def keyword_matches(keyword, text):
    words = [re.escape(w) for w in keyword.lower().split()]
    pattern = r"\b" + r"[\s\-]+".join(words) + r"\b"
    return re.search(pattern, text.lower()) is not None


def title_has_keyword(title, keywords):
    return any(keyword_matches(k, title) for k in keywords)


def new_funnel():
    return {
        "candidates": 0, "date_rejected": 0, "keyword_rejected": 0,
        "duplicate": 0, "collected": 0,
    }


def print_funnel(portal_name, f):
    print(
        f"[FUNNEL] {portal_name} | candidates={f['candidates']} | "
        f"collected={f['collected']} | "
        f"reject(date={f['date_rejected']}, "
        f"keyword={f['keyword_rejected']}, "
        f"duplicate={f['duplicate']})",
        flush=True,
    )


def extract_detik_date(date_elem):
    fallback = datetime.now().strftime("%Y-%m-%d")
    if not date_elem:
        return fallback
    raw = re.sub(r"\s+", " ", date_elem.get("title") or date_elem.text).strip()
    if "yang lalu" in raw:
        return fallback
    try:
        match = re.search(r"(\d{1,2})\s+([a-zA-Z]{3})\s+(\d{4})", raw)
        if match:
            day = f"{int(match.group(1)):02d}"
            mon = MONTH_MAP.get(match.group(2).lower()[:3], "01")
            year = match.group(3)
            return f"{year}-{mon}-{day}"
    except Exception:
        pass
    return fallback


def clean_title(title):
    if not title:
        return None
    title = re.sub(r"\s+", " ", title).strip()
    title = re.sub(r"\s+\d{6,}$", "", title)
    title = re.sub(r"^(FOTO:|VIDEO:|INFOGRAFIS:|BREAKING NEWS:|Ilustrasi|Gambar|Kumpulan)\s*[:\-]?\s*", "", title, flags=re.I)
    return title if len(title) > 15 else None


def within_window(dstr):
    try:
        y, m, d = map(int, dstr.split("-"))
        return MIN_DATE <= date(y, m, d) <= MAX_DATE
    except Exception:
        return False


def format_excel(file_path):
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=10)
    fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    grid_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    for cell in ws[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = font_body
            cell.border = grid_border
            cell.alignment = Alignment(horizontal="center" if cell.column in [1, 3] else "left", vertical="center")
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 65 if col_letter in ["D", "E"] else max(max_len + 3, 14)
    wb.save(file_path)


def new_portal_state():
    return {"collected": 0, "last_date": MAX_DATE, "exhausted": False}


# ==========================================
# DETIK  (original logs preserved)
# ==========================================
async def scrape_detik_batch(cfg, articles, seen_urls, seen_titles, state, quota, bar, funnel):
    if state["exhausted"] or quota <= 0:
        return 0

    portal_name = "Detik"
    f = funnel
    keywords = cfg["keywords"]
    host = cfg["detik_host"]
    index_base = cfg["detik_base"]
    origin = "https://" + host
    added = 0
    current = state["last_date"]

    print(f"\n[DETIK] Starting (index page, paginated, new-client-per-page)...", flush=True)

    while current >= MIN_DATE and added < quota and len(articles) < TARGET:
        date_iso = current.strftime("%Y-%m-%d")
        if cfg.get("detik_date_iso"):
            date_param = current.strftime("%Y-%m-%d")
        else:
            date_param = current.strftime("%m/%d/%Y")

        seen_this_date = set()
        prev_page_urls = None
        page = 0

        while True:
            page += 1
            if added >= quota or len(articles) >= TARGET:
                break

            print(f"[INDEX] {portal_name} | {index_base}?page={page}&date={date_param}", flush=True)
            headers = {
                "User-Agent": random.choice(USER_AGENTS),
                "Accept-Language": "id,en;q=0.9",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            }

            try:
                await asyncio.sleep(random.uniform(1.2, 2.5))
                async with httpx.AsyncClient(timeout=20.0, follow_redirects=True) as client:
                    res = await client.get(index_base, params={"page": page, "date": date_param}, headers=headers)

                if res.status_code != 200:
                    print(f"[INDEX] ERROR: {portal_name} | page={page} | HTTP {res.status_code}", flush=True)
                    break

                soup = BeautifulSoup(res.text, "html.parser")
                n_list_content = len(soup.select("div.list-content__item"))
                n_article = len(soup.select("article"))
                n_d_links = len(soup.select('a[href*="/d-"]'))
                print(
                    f"[DETIK-DEBUG] page={page} len={len(res.text)} "
                    f"list-content={n_list_content} "
                    f"article={n_article} "
                    f"d-links={n_d_links}",
                    flush=True,
                )

                links = (
                    soup.select("div.list-content__item a[href*='/d-']")
                    or soup.select("article a[href*='/d-']")
                    or soup.select(".media__text a[href*='/d-']")
                    or soup.select(".list-content a[href*='/d-']")
                )
                if not links:
                    links = soup.select('a[href*="/d-"]')
                    print(f"[DETIK-DEBUG] fell back to broad selector, count={len(links)}", flush=True)

                page_items = {}
                for link in links:
                    raw_href = link.get("href")
                    if not raw_href:
                        continue
                    art_url = urllib.parse.urljoin(origin, raw_href).split("?")[0].strip()
                    if host not in art_url:
                        continue
                    txt = clean_title(re.sub(r"\s+", " ", link.get_text()).strip()) or ""
                    if art_url not in page_items or len(txt) > len(page_items[art_url]):
                        page_items[art_url] = txt

                page_urls = set(page_items)

                if not page_urls:
                    print(f"[INDEX] {portal_name} | Date={date_iso} | page={page} | empty → next date", flush=True)
                    break

                new_for_date = page_urls - seen_this_date
                if not new_for_date:
                    print(f"[INDEX] {portal_name} | Date={date_iso} | page={page} | no unseen → next date", flush=True)
                    break

                if page_urls == prev_page_urls and prev_page_urls is not None:
                    print(f"[INDEX] {portal_name} | Date={date_iso} | page={page} | exact repeat → next date", flush=True)
                    break

                prev_page_urls = page_urls
                seen_this_date |= page_urls

                new = 0
                raw_links = len(page_items)
                for art_url, title_text in page_items.items():
                    if added >= quota or len(articles) >= TARGET:
                        break
                    f["candidates"] += 1
                    if not title_text:
                        continue
                    if art_url in seen_urls or title_text.lower() in seen_titles:
                        f["duplicate"] += 1
                        continue
                    m = re.search(r"/(\d{4})/(\d{2})/(\d{2})", art_url)
                    dstr = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else date_iso
                    if not within_window(dstr):
                        f["date_rejected"] += 1
                        continue
                    if not title_has_keyword(title_text, keywords):
                        f["keyword_rejected"] += 1
                        continue
                    seen_urls.add(art_url)
                    seen_titles.add(title_text.lower())
                    articles.append({
                        "id": 0, "source": "Detik", "date": dstr,
                        "title": title_text, "label": None, "url": art_url,
                    })
                    f["collected"] += 1
                    added += 1
                    new += 1
                    if bar is not None:
                        bar.update(1)

                print(
                    f"[INDEX] {portal_name} | Date={date_iso} | page={page} | "
                    f"raw_links={raw_links} | new={new} | Collected={len(articles)}/{TARGET}",
                    flush=True,
                )

            except Exception as e:
                print(f"[DETIK] error {date_iso} page={page}: {e}", flush=True)
                break

        state["last_date"] = current
        current -= timedelta(days=1)

    if current < MIN_DATE:
        state["exhausted"] = True
    state["collected"] += added
    return added


# ==========================================
# CNN  (original logs preserved)
# ==========================================
def scrape_cnn_batch(cfg, articles, seen_urls, seen_titles, state, quota, bar, funnel):
    if state["exhausted"] or quota <= 0:
        return 0

    portal_name = "CNN Indonesia"
    f = funnel
    keywords = cfg["keywords"]
    section = cfg["cnn_section"]
    base = f"https://www.cnnindonesia.com/{cfg['cnn_indeks']}"
    added = 0
    current = state["last_date"]

    print(f"\n[CNN] Starting (date+page method)...", flush=True)

    while current >= MIN_DATE and added < quota and len(articles) < TARGET:
        date_slash = current.strftime("%Y/%m/%d")
        date_iso = current.strftime("%Y-%m-%d")
        page = 1
        empty_streak = 0

        while added < quota and len(articles) < TARGET and empty_streak < 2:
            url = f"{base}?date={date_slash}&page={page}"
            print(f"[INDEX] {portal_name} | Date={date_iso} | Page={page}", flush=True)
            headers = {"User-Agent": random.choice(USER_AGENTS), "Accept-Language": "id,en;q=0.9"}
            try:
                time.sleep(random.uniform(1.0, 2.0))
                res = requests.get(url, headers=headers, timeout=12)
                if res.status_code != 200 or "gokuProps" in (res.text or "")[:1500]:
                    print(f"[INDEX] {portal_name} | Date={date_iso} Page={page} | blocked/empty (HTTP {res.status_code})", flush=True)
                    empty_streak += 1
                    page += 1
                    continue

                soup = BeautifulSoup(res.content, "html.parser")
                article_nodes = soup.select("article")
                if not article_nodes:
                    print(
                        f"[INDEX] {portal_name} | Date={date_iso} Page={page} | "
                        f"raw_links=0 | new=0 | "
                        f"Collected={len(articles)}/{TARGET} | Remaining={TARGET - len(articles)}",
                        flush=True,
                    )
                    break

                seen_in_page = set()
                candidate_anchors = []
                for node in article_nodes:
                    for a in node.find_all("a", href=True):
                        h = a["href"].split("?")[0]
                        if h not in seen_in_page:
                            seen_in_page.add(h)
                            candidate_anchors.append(a)

                new = 0
                raw_links = 0
                for a in candidate_anchors:
                    if added >= quota or len(articles) >= TARGET:
                        break
                    href = a["href"].split("?")[0]
                    if section not in href:
                        continue
                    raw_links += 1
                    f["candidates"] += 1
                    title_text = clean_title(re.sub(r"\s+", " ", a.get_text()).strip())
                    if not title_text:
                        continue
                    if href in seen_urls or title_text.lower() in seen_titles:
                        f["duplicate"] += 1
                        continue
                    m = re.search(r"/(\d{4})/(\d{2})/(\d{2})", href)
                    dstr = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else date_iso
                    if not within_window(dstr):
                        f["date_rejected"] += 1
                        continue
                    if not title_has_keyword(title_text, keywords):
                        f["keyword_rejected"] += 1
                        continue
                    seen_urls.add(href)
                    seen_titles.add(title_text.lower())
                    articles.append({
                        "id": 0, "source": "CNN Indonesia", "date": dstr,
                        "title": title_text, "label": None, "url": href,
                    })
                    f["collected"] += 1
                    added += 1
                    new += 1
                    if bar is not None:
                        bar.update(1)

                print(
                    f"[INDEX] {portal_name} | Date={date_iso} Page={page} | "
                    f"raw_links={raw_links} | new={new} | "
                    f"Collected={len(articles)}/{TARGET} | Remaining={TARGET - len(articles)}",
                    flush=True,
                )
                if raw_links == 0:
                    break
                page += 1

            except Exception as e:
                print(f"[CNN] error: {e}", flush=True)
                break

        state["last_date"] = current
        current -= timedelta(days=1)

    if current < MIN_DATE:
        state["exhausted"] = True
    state["collected"] += added
    return added


# ==========================================
# KOMPAS (original logs preserved)
# ==========================================
async def scrape_kompas_batch(cfg, articles, seen_urls, seen_titles, state, quota, bar, funnel):
    if state["exhausted"] or quota <= 0:
        return 0

    portal_name = "Kompas"
    f = funnel
    keywords = cfg["keywords"]
    site = cfg["kompas_site"]
    patterns = cfg.get("kompas_url_patterns") or [f"{cfg['kompas_host']}/read/"]
    added = 0
    current = state["last_date"]
    fails = 0

    print(f"\n[KOMPAS] Starting with Playwright (async)...", flush=True)

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            context = await browser.new_context(
                user_agent=random.choice(USER_AGENTS),
                locale="id-ID",
                viewport={"width": 1366, "height": 768},
            )
            page = await context.new_page()

            while current >= MIN_DATE and added < quota and len(articles) < TARGET and fails < 5:
                date_iso = current.strftime("%Y-%m-%d")
                page_num = 1

                while added < quota and len(articles) < TARGET and fails < 5:
                    url = f"https://indeks.kompas.com/?site={site}&date={date_iso}&page={page_num}"
                    print(f"[INDEX] {portal_name} | Date={date_iso} | Page={page_num}", flush=True)

                    try:
                        await page.goto(url, timeout=30000, wait_until="domcontentloaded")
                        await asyncio.sleep(random.uniform(2, 3))

                        prev = 0
                        for _ in range(6):
                            await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                            await asyncio.sleep(1.2)
                            cur = await page.evaluate('document.querySelectorAll(\'a[href*="/read/"]\').length')
                            if cur == prev:
                                break
                            prev = cur
                        await page.evaluate("window.scrollTo(0, 0)")
                        await asyncio.sleep(0.4)

                        content = await page.content()
                        if "gokuProps" in content or len(content) < 2000:
                            fails += 1
                            print(f"[KOMPAS] blocked {date_iso} p{page_num} ({fails}/5)", flush=True)
                            break

                        fails = 0
                        css_sel = ", ".join(f'a[href*="{p}"]' for p in patterns)
                        links = await page.query_selector_all(css_sel)

                        new = 0
                        raw_links = 0
                        seen_in_page = set()
                        for link in links:
                            if added >= quota or len(articles) >= TARGET:
                                break
                            href = await link.get_attribute("href")
                            if not href:
                                continue
                            href = urllib.parse.urljoin("https://indeks.kompas.com/", href).split("?")[0]
                            if href in seen_in_page:
                                continue
                            seen_in_page.add(href)
                            if not any(p in href for p in patterns):
                                continue
                            raw_links += 1
                            f["candidates"] += 1
                            if href in seen_urls:
                                f["duplicate"] += 1
                                continue
                            title_text = clean_title((await link.inner_text()).strip())
                            if not title_text:
                                continue
                            if title_text.lower() in seen_titles:
                                f["duplicate"] += 1
                                continue
                            m = re.search(r"/(\d{4})/(\d{2})/(\d{2})", href)
                            if m:
                                dstr = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
                                if not within_window(dstr):
                                    f["date_rejected"] += 1
                                    continue
                            else:
                                dstr = date_iso
                            if not title_has_keyword(title_text, keywords):
                                f["keyword_rejected"] += 1
                                continue
                            seen_urls.add(href)
                            seen_titles.add(title_text.lower())
                            articles.append({
                                "id": 0, "source": "Kompas", "date": dstr,
                                "title": title_text, "label": None, "url": href,
                            })
                            f["collected"] += 1
                            added += 1
                            new += 1
                            if bar is not None:
                                bar.update(1)

                        print(
                            f"[INDEX] {portal_name} | Date={date_iso} Page={page_num} | "
                            f"raw_links={raw_links} | new={new} | "
                            f"Collected={len(articles)}/{TARGET} | Remaining={TARGET - len(articles)}",
                            flush=True,
                        )
                        if raw_links == 0:
                            break
                        page_num += 1

                    except Exception as e:
                        print(f"[KOMPAS] error: {e}", flush=True)
                        fails += 1
                        break

                state["last_date"] = current
                current -= timedelta(days=1)

            await browser.close()
    except Exception as e:
        print(f"[KOMPAS] browser failed: {e}", flush=True)
        state["exhausted"] = True

    if current < MIN_DATE or fails >= 5:
        state["exhausted"] = True
    state["collected"] += added
    return added


# ==========================================
# TRIBUN (original logs preserved)
# ==========================================
async def tribun_pub_date(page, url):
    try:
        await page.goto(url, timeout=30000, wait_until="domcontentloaded")
        for sel in ('meta[property="article:published_time"]',
                    'meta[name="content_PublishedDate"]',
                    'meta[itemprop="datePublished"]'):
            el = await page.query_selector(sel)
            if el:
                v = (await el.get_attribute("content") or "").strip()
                m = re.match(r"(\d{4})-(\d{2})-(\d{2})", v)
                if m:
                    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
        html = await page.content()
        m = re.search(r'"datePublished"\s*:\s*"(\d{4})-(\d{2})-(\d{2})', html)
        if m:
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    except Exception as e:
        print(f"[TRIBUN] date fetch error: {e}", flush=True)
    return None


async def scrape_tribun_batch(cfg, articles, seen_urls, seen_titles, state, quota, bar, funnel):
    if state["exhausted"] or quota <= 0:
        return 0

    portal_name = "Tribun News"
    f = funnel
    keywords = cfg["keywords"]
    marker = cfg["tribun_marker"]
    added = 0
    fails = 0

    print(f"\n[TRIBUN] Starting with Playwright (async)...", flush=True)

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            context = await browser.new_context(
                user_agent=random.choice(USER_AGENTS),
                locale="id-ID",
                viewport={"width": 1366, "height": 768},
            )
            page = await context.new_page()
            art_page = await context.new_page()

            for channel in cfg["tribun_channels"]:
                if added >= quota or len(articles) >= TARGET or fails >= 5:
                    break
                page_num = 1
                while added < quota and len(articles) < TARGET and fails < 5:
                    url = f"https://www.tribunnews.com/{channel}?page={page_num}"
                    print(f"[INDEX] {portal_name} | Channel={channel} | Page={page_num}", flush=True)

                    try:
                        await page.goto(url, timeout=30000, wait_until="domcontentloaded")
                        await asyncio.sleep(random.uniform(2, 3))
                        content = await page.content()
                        if "gokuProps" in content or len(content) < 2000:
                            fails += 1
                            print(f"[TRIBUN] blocked {channel} page {page_num} ({fails}/5)", flush=True)
                            page_num += 1
                            continue
                        fails = 0

                        ART_ID_RE = re.compile(r"^/[\w\-]+/\d+/")
                        links = await page.query_selector_all('a[href*="tribunnews.com"]')
                        seen_in_page = set()
                        new = 0
                        raw_links = 0

                        for link in links:
                            if added >= quota or len(articles) >= TARGET:
                                break
                            href = await link.get_attribute("href")
                            if not href:
                                continue
                            href = href.split("?")[0].split("#")[0]
                            if href in seen_in_page:
                                continue
                            seen_in_page.add(href)
                            if marker not in href:
                                continue
                            path = urllib.parse.urlparse(href).path
                            if not ART_ID_RE.match(path):
                                continue
                            raw_links += 1
                            f["candidates"] += 1
                            if href in seen_urls:
                                f["duplicate"] += 1
                                continue
                            title_text = clean_title((await link.inner_text()).strip())
                            if not title_text:
                                continue
                            if title_text.lower() in seen_titles:
                                f["duplicate"] += 1
                                continue
                            if not title_has_keyword(title_text, keywords):
                                f["keyword_rejected"] += 1
                                continue

                            await asyncio.sleep(random.uniform(1.0, 2.0))
                            dstr = await tribun_pub_date(art_page, href)
                            if dstr is None or not within_window(dstr):
                                f["date_rejected"] += 1
                                continue

                            seen_urls.add(href)
                            seen_titles.add(title_text.lower())
                            articles.append({
                                "id": 0, "source": "Tribun News", "date": dstr,
                                "title": title_text, "label": None, "url": href,
                            })
                            f["collected"] += 1
                            added += 1
                            new += 1
                            if bar is not None:
                                bar.update(1)

                        print(
                            f"[INDEX] {portal_name} | Channel={channel} Page={page_num} | "
                            f"raw_links={raw_links} | new={new} | "
                            f"Collected={len(articles)}/{TARGET} | Remaining={TARGET - len(articles)}",
                            flush=True,
                        )
                        if raw_links == 0:
                            break
                        page_num += 1

                    except Exception as e:
                        print(f"[TRIBUN] error: {e}", flush=True)
                        fails += 1
                        page_num += 1

            await browser.close()
    except Exception as e:
        print(f"[TRIBUN] browser failed: {e}", flush=True)
        state["exhausted"] = True

    if fails >= 5:
        state["exhausted"] = True
    state["collected"] += added
    return added


# ==========================================
# RUN ONE DOMAIN (backfill controller)
# ==========================================
async def run_domain(domain_key):
    cfg = DOMAINS[domain_key]
    articles = []
    seen_urls = set()
    seen_titles = set()
    funnels = {name: new_funnel() for name in PORTAL_ORDER}
    states = {name: new_portal_state() for name in PORTAL_ORDER}

    print(f"\n========== Starting {cfg['label']} scraper (Target: {TARGET}) ==========", flush=True)
    bar = tqdm(total=TARGET, desc=f"Harvesting {cfg['label']}", unit="rows")

    scrapers = {
        "Detik": scrape_detik_batch,
        "Kompas": scrape_kompas_batch,
        "CNN Indonesia": scrape_cnn_batch,
        "Tribun News": scrape_tribun_batch,
    }

    max_rounds = 8
    for round_i in range(1, max_rounds + 1):
        if len(articles) >= TARGET:
            break
        made_progress = False

        for portal in PORTAL_ORDER:
            if len(articles) >= TARGET:
                break
            st = states[portal]
            if st["exhausted"] or st["collected"] >= PORTAL_CAP:
                continue

            remaining_global = TARGET - len(articles)
            remaining_portal = PORTAL_CAP - st["collected"]
            quota = min(remaining_global, remaining_portal, 80)
            if quota <= 0:
                continue

            if portal == "CNN Indonesia":
                added = scrape_cnn_batch(cfg, articles, seen_urls, seen_titles, st, quota, bar, funnels[portal])
            else:
                added = await scrapers[portal](cfg, articles, seen_urls, seen_titles, st, quota, bar, funnels[portal])

            if added > 0:
                made_progress = True

        if not made_progress:
            break

    bar.close()

    print(f"\n[FUNNEL] {cfg['label']} drop-off by portal:", flush=True)
    for name in PORTAL_ORDER:
        print_funnel(name, funnels[name])

    articles = articles[:TARGET]
    for i, row in enumerate(articles, 1):
        row["id"] = i

    os.makedirs("data", exist_ok=True)
    df = pd.DataFrame(articles, columns=OUTPUT_COLUMNS)
    df["label"] = pd.NA
    df.to_csv(cfg["output_file"], index=False)
    print(f"[SUCCESS] Saved {len(df)} rows → {cfg['output_file']}", flush=True)

    xlsx = cfg["output_file"].replace(".csv", ".xlsx")
    df.to_excel(xlsx, index=False)
    try:
        format_excel(xlsx)
        print(f"[SUCCESS] Excel also saved → {xlsx}", flush=True)
    except Exception:
        pass
    return df 


# ==========================================
# ENTRY
# ==========================================
async def run_all():
    for key in DOMAINS:
        await run_domain(key)


# Example:
# await run_domain("technology")
# await run_domain("sport")
await run_all()

# if __name__ == "__main__":
#     asyncio.run(run_domain("technology"))